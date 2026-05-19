import io
import logging
import mimetypes
import os
import zipfile
from dataclasses import dataclass
from typing import Iterable, List, Optional

import fitz
from xml.etree import ElementTree

from app.models.base import KBDocument

logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 1_048_576  # 1 MB
ALLOWED_EXTENSIONS = {
    # Documents & text
    ".pdf",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".js",
    ".html",
    ".xml",
    ".rtf",
    ".doc",
    ".docx",
    ".odt",
    ".xls",
    ".xlsx",
    ".ods",
    # Images
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".gif",
    ".heic",
    ".heif",
    # Audio / Video passthrough
    ".mp4",
    ".flv",
    ".mov",
    ".mpeg",
    ".mpg",
    ".webm",
    ".wmv",
    ".3gp",
    ".mp3",
    ".wav",
    ".m4a",
    ".aac",
    ".flac",
    ".ogg",
    ".pcm",
}
GEMINI_DIRECT_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".gif"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}


@dataclass
class KBUploadPayload:
    filename: str
    content: bytes
    mime_type: str


class KBProcessorService:
    """Shared service for validating KB uploads and extracting fallback text."""

    @staticmethod
    def infer_mime_type(filename: str, fallback: Optional[str] = None) -> str:
        guessed, _ = mimetypes.guess_type(filename)
        return fallback or guessed or "application/octet-stream"

    @classmethod
    def validate_upload(
        cls,
        filename: str,
        content_bytes: bytes,
        mime_type: Optional[str] = None,
    ) -> KBUploadPayload:
        ext = os.path.splitext(filename)[1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
            raise ValueError(f"Unsupported file type {ext or '(none)'}. Allowed: {allowed}")
        if len(content_bytes) > MAX_FILE_SIZE:
            raise ValueError(
                f"File exceeds 1 MB limit ({len(content_bytes) / 1_048_576:.1f} MB). Please upload a smaller file."
            )
        return KBUploadPayload(
            filename=filename,
            content=content_bytes,
            mime_type=cls.infer_mime_type(filename, mime_type),
        )

    @staticmethod
    async def extract_text(filename: str, content_bytes: bytes) -> str:
        """Extract text from text-like docs and provide placeholders for image files."""
        ext = os.path.splitext(filename)[1].lower()

        if ext in {".txt", ".md", ".csv", ".json", ".js", ".html", ".xml", ".rtf"}:
            return content_bytes.decode("utf-8", errors="ignore")

        if ext == ".pdf":
            try:
                doc = fitz.open("pdf", content_bytes)
                return "\n".join(page.get_text() for page in doc)
            except Exception as exc:
                return f"[Error extracting PDF: {exc}]"

        if ext == ".docx":
            return KBProcessorService._extract_docx_plain(content_bytes)

        if ext == ".doc":
            return "[Doc file detected. Preview limited. Please convert to DOCX for full extraction.]"

        if ext in {".xlsx", ".xls"}:
            try:
                import openpyxl

                wb = openpyxl.load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
                sheet = wb.active
                rows = []
                for idx, row in enumerate(sheet.iter_rows(values_only=True)):
                    if idx >= 50:  # cap preview
                        break
                    rows.append(", ".join("" if v is None else str(v) for v in row[:20]))
                return "\n".join(rows) or "[Empty spreadsheet]"
            except ImportError:
                # Fallback to lightweight XML parse of first sheet (XLSX) or xlrd hint (XLS)
                if ext == ".xlsx":
                    return KBProcessorService._extract_xlsx_plain(content_bytes)
                if ext == ".xls":
                    return "[XLS file uploaded. Preview unavailable; download or convert to XLSX/CSV for inline view.]"
                return "[Spreadsheet preview limited; install openpyxl for full extract.]"
            except Exception as exc:
                return f"[Error extracting spreadsheet: {exc}]"

        if ext in {".ods"}:
            return "[ODS file uploaded. Inline text preview limited; download to view.]"

        if ext in IMAGE_EXTENSIONS:
            return f"[Image context available for Gemini: {filename}]"

        if ext in {".mp4", ".flv", ".mov", ".mpeg", ".mpg", ".webm", ".wmv", ".3gp", ".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".pcm"}:
            return f"[Media file uploaded: {filename}]"

        return f"[Unsupported file type: {ext}]"

    @staticmethod
    def _extract_docx_plain(content_bytes: bytes) -> str:
        """Lightweight DOCX text extraction without python-docx."""
        try:
            with zipfile.ZipFile(io.BytesIO(content_bytes)) as zf:
                if "word/document.xml" not in zf.namelist():
                    return "[DOCX missing document.xml]"
                xml_bytes = zf.read("word/document.xml")
            tree = ElementTree.fromstring(xml_bytes)
            # Namespace handling
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = []
            for p in tree.findall(".//w:p", ns):
                texts = [t.text for t in p.findall(".//w:t", ns) if t.text]
                if texts:
                    paragraphs.append("".join(texts))
            return "\n".join(paragraphs) or "[DOCX contained no text content]"
        except Exception as exc:
            return f"[Error extracting DOCX: {exc}]"

    @staticmethod
    def _extract_xlsx_plain(content_bytes: bytes) -> str:
        """Lightweight XLSX extractor without openpyxl — grabs first 50 rows/20 cols."""
        try:
            with zipfile.ZipFile(io.BytesIO(content_bytes)) as zf:
                shared_strings = []
                if "xl/sharedStrings.xml" in zf.namelist():
                    ss_tree = ElementTree.fromstring(zf.read("xl/sharedStrings.xml"))
                    ns = {"t": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    for si in ss_tree.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si"):
                        texts = []
                        for t in si.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"):
                            if t.text:
                                texts.append(t.text)
                        shared_strings.append("".join(texts))

                sheet_name = next((n for n in zf.namelist() if n.startswith("xl/worksheets/sheet1")), None)
                if not sheet_name:
                    return "[XLSX contains no sheet1 to preview]"

                sheet_tree = ElementTree.fromstring(zf.read(sheet_name))
                rows_out = []
                ns_main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

                for r_idx, row in enumerate(sheet_tree.findall(f".//{ns_main}row")):
                    if r_idx >= 50:
                        break
                    row_vals = []
                    for c_idx, cell in enumerate(row.findall(f"{ns_main}c")):
                        if c_idx >= 20:
                            break
                        cell_type = cell.get("t")
                        v = cell.find(f"{ns_main}v")
                        val_text = v.text if v is not None else ""
                        if cell_type == "s":
                            try:
                                val_text = shared_strings[int(val_text)]
                            except Exception:
                                pass
                        row_vals.append(val_text)
                    rows_out.append(", ".join(row_vals))

                return "\n".join(rows_out) or "[Empty sheet]"
        except Exception as exc:
            return f"[Error extracting XLSX: {exc}]"

    @staticmethod
    def _extract_xls_plain(content_bytes: bytes) -> str:
        """Deprecated XLS extractor (not used); kept for reference."""
        return "[XLS preview unavailable]"

    @classmethod
    async def build_kb_document(
        cls,
        filename: str,
        content_bytes: bytes,
        mime_type: Optional[str] = None,
    ) -> KBDocument:
        upload = cls.validate_upload(filename, content_bytes, mime_type)
        text = await cls.extract_text(upload.filename, upload.content)
        return KBDocument(
            filename=upload.filename,
            content=text,
            file_type=upload.mime_type,
            size_bytes=len(upload.content),
        )

    @classmethod
    def should_attach_to_gemini(cls, filename: str, mime_type: Optional[str] = None) -> bool:
        ext = os.path.splitext(filename)[1].lower()
        resolved_mime = cls.infer_mime_type(filename, mime_type)
        return ext in GEMINI_DIRECT_EXTENSIONS or resolved_mime.startswith("image/") or resolved_mime == "application/pdf"

    @classmethod
    def build_gemini_attachment(
        cls,
        filename: str,
        content_bytes: bytes,
        mime_type: Optional[str] = None,
    ) -> Optional[dict]:
        upload = cls.validate_upload(filename, content_bytes, mime_type)
        if not cls.should_attach_to_gemini(upload.filename, upload.mime_type):
            logger.info(
                "KB_GEMINI_ATTACHMENT_SKIPPED filename=%s mime_type=%s reason=unsupported_for_direct_file_param",
                upload.filename,
                upload.mime_type,
            )
            return None
        logger.info(
            "KB_GEMINI_ATTACHMENT_READY filename=%s mime_type=%s size_bytes=%d",
            upload.filename,
            upload.mime_type,
            len(upload.content),
        )
        return {
            "filename": upload.filename,
            "mime_type": upload.mime_type,
            "data": upload.content,
        }

    @classmethod
    async def process_uploads(cls, files: Iterable[tuple]) -> List[KBDocument]:
        """Process tuples of (filename, content_bytes[, mime_type])."""
        documents: List[KBDocument] = []
        for item in files:
            if len(item) == 2:
                filename, content = item
                mime_type = None
            else:
                filename, content, mime_type = item
            documents.append(await cls.build_kb_document(filename, content, mime_type))
        return documents
